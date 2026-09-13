#!/usr/bin/env python3
"""Build the English hand-off report for the committed RQ4 figure."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_FIGURES = ROOT / "paper" / "figures"


def set_cell(cell, fill: str | None = None) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
    if fill:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def labelled(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(label).bold = True
    paragraph.add_run(text)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--figures", type=Path, default=DEFAULT_FIGURES)
    args = parser.parse_args()
    with (args.figures / "Fig1-summary.csv").open(newline="", encoding="utf-8") as source:
        summary = list(csv.DictReader(source))

    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(9.5)
    for section in document.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
    for style_name in ("Title", "Heading 1"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Title"].font.size = Pt(18)
    document.styles["Heading 1"].font.size = Pt(12)
    title_style_properties = document.styles["Title"]._element.get_or_add_pPr()
    title_style_border = title_style_properties.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_properties.remove(title_style_border)
    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(3)
    title.add_run("Joggle figure report")
    title_properties = title._p.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    target = document.add_paragraph("Target  EuroSys 2027 fall cycle working manuscript")
    target.paragraph_format.space_after = Pt(7)
    heading = document.add_heading("Figure 1  Structural block policy tradeoff", level=1)
    heading.paragraph_format.space_before = Pt(3)
    heading.paragraph_format.space_after = Pt(3)
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_after = Pt(5)
    picture.add_run().add_picture(str(args.figures / "Fig1.png"), width=Inches(6.7))
    labelled(
        document,
        "Claim  ",
        "The same operator-independent composition of split, reorder, and "
        "scalar promotion improves every paired call in three model pilots, while "
        "substantially increasing generated C source size.",
    )
    labelled(
        document,
        "Caption  ",
        "Descriptive same-process pilot results for a source-defined "
        "function-body block policy. (A) Each point is one adjacent baseline/candidate "
        "timed pair; horizontal bars are medians and vertical lines are interquartile "
        "ranges. The dashed line denotes equal latency. Observations are repeated "
        "technical calls within one unisolated process (n = 40, 20, and 20), not "
        "independent experimental units; no inferential test is reported. Candidate "
        "and baseline outputs are bit-identical for every pair. (B) Exact generated C "
        "source-byte ratios from the corresponding preserved artifact records; the "
        "dashed line denotes equal size. These pilots do not compare Joggle with a "
        "production runtime.",
    )
    labelled(
        document,
        "Citation location  ",
        "Section 5, RQ4 paragraph beginning ‘A later same-process diagnostic selects "
        "factors four and seven…’.",
    )
    labelled(
        document,
        "Reproducibility  ",
        "paper/figures/block_tradeoff.py reads the three committed *-block-pilot.csv "
        "timing files and paper/data/block-artifact-pilot.csv.",
    )
    values_heading = document.add_heading("Derived values", level=1)
    values_heading.paragraph_format.space_before = Pt(5)
    values_heading.paragraph_format.space_after = Pt(3)
    table = document.add_table(rows=1, cols=5)
    headers = ["Model", "n", "Median ratio", "IQR", "C bytes ratio"]
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        cell.text = value
        set_cell(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            run.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in summary:
        cells = table.add_row().cells
        values = [
            row["model"],
            row["observations"],
            f"{float(row['paired_ratio_median']):.3f}",
            f"[{float(row['paired_ratio_q1']):.3f}, "
            f"{float(row['paired_ratio_q3']):.3f}]",
            f"{float(row['c_source_ratio']):.3f}",
        ]
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
            set_cell(cell)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.save(args.figures / "Fig1-report.docx")


if __name__ == "__main__":
    main()
